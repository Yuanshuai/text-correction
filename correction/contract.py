"""合同模板解析：提取占位符字段及上下文。"""
import re
from dataclasses import dataclass
from docx import Document


@dataclass
class PlaceholderField:
    field_name: str
    field_type: str  # "text" | "number" | "select" | "percent" | "date" | "amount"
    paragraph_index: int
    context_before: str
    context_after: str
    full_paragraph: str
    table_index: int | None = None
    table_row: int | None = None
    table_col: int | None = None


def _detect_field_type(text: str, field_name: str) -> str:
    combined = field_name + text
    if re.search(r'第\s*种', combined):
        return "select"
    if re.search(r'(日|天|工作日|小时)', text) and re.search(r'(超过|满|届满|之前|内)', combined):
        return "number"
    if re.search(r'[%％]', text) or re.search(r'(比例|比例|率)', field_name):
        return "percent"
    if re.search(r'(年.*月.*日|月.*日|签订日期)', combined):
        return "date"
    if re.search(r'(元|大写|小写|价款|金额|费用)', combined):
        return "amount"
    if re.search(r'(人次|份|条|项|款)', text):
        return "number"
    return "text"


def _find_placeholders_in_paragraph(text: str, para_index: int) -> list[PlaceholderField]:
    results = []
    # Pattern 1: "字段名：     " or "字段名：" with trailing spaces
    for m in re.finditer(r'([^：:\n]{1,30})[：:]\s{2,}', text):
        field_name = m.group(1).strip()
        field_name = re.sub(r'^[\d.]+\s*', '', field_name)
        if not field_name or len(field_name) < 2:
            continue
        results.append(PlaceholderField(
            field_name=field_name,
            field_type=_detect_field_type("", field_name),
            paragraph_index=para_index,
            context_before=text[:m.start()],
            context_after=text[m.end():],
            full_paragraph=text,
        ))

    # Pattern 2: "第种" (select placeholder)
    for m in re.finditer(r'第\s*种', text):
        start = max(0, m.start() - 20)
        results.append(PlaceholderField(
            field_name="选择项",
            field_type="select",
            paragraph_index=para_index,
            context_before=text[start:m.start()],
            context_after=text[m.end():],
            full_paragraph=text,
        ))

    # Pattern 3: "超过/满/内     日/天" (number placeholder)
    for m in re.finditer(r'(?:超过|满|内|届满)\s{2,}(?:日|天|工作日|小时)', text):
        field_name = m.group(0).strip()
        results.append(PlaceholderField(
            field_name=field_name,
            field_type="number",
            paragraph_index=para_index,
            context_before=text[:m.start()],
            context_after=text[m.end():],
            full_paragraph=text,
        ))

    # Pattern 4: "的% " or "的 ％ " (percent placeholder)
    for m in re.finditer(r'(?:的|为)\s*(?:[%％])', text):
        results.append(PlaceholderField(
            field_name="百分比",
            field_type="percent",
            paragraph_index=para_index,
            context_before=text[:m.start()],
            context_after=text[m.end():],
            full_paragraph=text,
        ))

    # Pattern 5: inline blanks in sentences "乙方指派     作为"
    for m in re.finditer(r'([\u4e00-\u9fff]{2,8})\s{3,}([\u4e00-\u9fff]{2,8})', text):
        field_name = f"{m.group(1)}___{m.group(2)}"
        results.append(PlaceholderField(
            field_name=field_name,
            field_type="text",
            paragraph_index=para_index,
            context_before=text[:m.start()],
            context_after=text[m.end():],
            full_paragraph=text,
        ))

    return results


def parse_contract(path: str) -> tuple[list[PlaceholderField], list[str]]:
    """解析合同 DOCX，返回 (占位符字段列表, 全部段落文本列表)。"""
    doc = Document(path)
    all_paragraphs = [p.text for p in doc.paragraphs]
    fields: list[PlaceholderField] = []

    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue
        fields.extend(_find_placeholders_in_paragraph(text, i))

    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                text = cell.text.strip()
                if not text:
                    header = table.rows[0].cells[ci].text.strip() if ri > 0 and ci < len(table.rows[0].cells) else f"列{ci}"
                    fields.append(PlaceholderField(
                        field_name=header,
                        field_type="text",
                        paragraph_index=-1,
                        context_before="",
                        context_after="",
                        full_paragraph=f"[表格{ti} 行{ri}] {header}: (空)",
                        table_index=ti,
                        table_row=ri,
                        table_col=ci,
                    ))

    return fields, all_paragraphs
